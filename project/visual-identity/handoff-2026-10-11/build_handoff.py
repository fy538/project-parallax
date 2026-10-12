"""Build the agent handoff from its Markdown companion using python-docx."""
from pathlib import Path
import re
from urllib.parse import urlparse
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image

HERE = Path(__file__).resolve().parent
SOURCE = HERE / 'parallax-visual-research-handoff.md'
OUTPUT = HERE / 'parallax-visual-research-handoff.docx'
doc = Document()
# Strip any decorative paragraph rules inherited from the runtime template.
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.72), Inches(.7)
sec.left_margin = sec.right_margin = Inches(.78)
sec.header_distance = sec.footer_distance = Inches(.3)
WIDTH = 6.94

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number', 'Caption']:
    st = doc.styles[name]
    st.font.name = 'Arial'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.line_spacing = 1.12
    st.paragraph_format.widow_control = True
doc.styles['Title'].font.size = Pt(27)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(12)
for name, size, before, after in [('Heading 1',17,17,9), ('Heading 2',12.5,11,5), ('Heading 3',11.5,9,4)]:
    st = doc.styles[name]
    st.font.size, st.font.bold = Pt(size), True
    st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(before), Pt(after)
    st.paragraph_format.keep_with_next = True
doc.styles['Caption'].font.size = Pt(9)
doc.styles['Caption'].font.italic = False
doc.styles['Caption'].paragraph_format.line_spacing = 1.05
doc.styles['Caption'].paragraph_format.space_after = Pt(9)
doc.styles['List Bullet'].paragraph_format.left_indent = Inches(.18)
doc.styles['List Bullet'].paragraph_format.first_line_indent = Inches(-.12)

def field(p, name):
    r = p.add_run()
    f = OxmlElement('w:fldSimple')
    f.set(qn('w:instr'), name)
    r._r.addnext(f)

hp = sec.header.paragraphs[0]
hp.text = 'PARALLAX  /  VISUAL RESEARCH HANDOFF'
hp.style = doc.styles['Normal']
hp.runs[0].font.size = Pt(8)
hp.paragraph_format.space_after = Pt(0)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('October 10 and 11 2026   |   ').font.size = Pt(8)
field(fp,'PAGE')
fp.add_run(' of ').font.size = Pt(8)
field(fp,'NUMPAGES')
doc.core_properties.title = 'Parallax visual research handoff'
doc.core_properties.subject = 'October 10 and 11 2026 creative and production research'
doc.core_properties.author = 'Parallax / Codex'
doc.core_properties.keywords = 'visual identity, constructivism, surrealism, featureless faces, Remotion, research handoff'

bookmark_id = 0
def bookmark(p, name):
    global bookmark_id
    bookmark_id += 1
    a, b = OxmlElement('w:bookmarkStart'), OxmlElement('w:bookmarkEnd')
    a.set(qn('w:id'),str(bookmark_id)); a.set(qn('w:name'),name)
    b.set(qn('w:id'),str(bookmark_id))
    p._p.insert(0,a); p._p.append(b)

def hyperlink(p, label, target, internal=False):
    h = OxmlElement('w:hyperlink')
    if internal:
        h.set(qn('w:anchor'),target)
    else:
        if not urlparse(target).scheme:
            target = (HERE / target).resolve().as_uri()
        rid = p.part.relate_to(target,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
        h.set(qn('r:id'),rid)
    r, props = OxmlElement('w:r'), OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'),'234F68' if not internal else '000000')
    props.append(color)
    if not internal:
        u = OxmlElement('w:u'); u.set(qn('w:val'),'single');props.append(u)
    r.append(props)
    t = OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def inline(p, value, table_head=False):
    # Balanced URL parentheses preserve Commons source links.
    i = 0
    while i < len(value):
        if value.startswith('**',i):
            end=value.find('**',i+2)
            if end>=0:
                r=p.add_run(value[i+2:end]);r.bold=True;i=end+2;continue
        if value[i]=='`':
            end=value.find('`',i+1)
            if end>=0:
                r=p.add_run(value[i+1:end]);r.font.name='Courier New';r.font.size=Pt(9)
                i=end+1;continue
        if value[i]=='[':
            split=value.find('](',i+1)
            if split>=0:
                depth=1;j=split+2
                while j<len(value) and depth:
                    if value[j]=='(':depth+=1
                    if value[j]==')':depth-=1
                    j+=1
                if depth==0:
                    hyperlink(p,value[i+1:split],value[split+2:j-1]);i=j;continue
        j=i+1
        while j<len(value) and value[j] not in '`[' and not value.startswith('**',j):j+=1
        p.add_run(value[i:j]);i=j
    if table_head:
        for r in p.runs:r.font.color.rgb=RGBColor(255,255,255);r.bold=True

def add_table(lines):
    rows=[[x.strip() for x in line.strip().strip('|').split('|')] for line in lines]
    rows=[r for r in rows if not all(re.match(r'^:?-+:?$',x) for x in r)]
    table=doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    widths=[1.3,2.85,2.79] if len(rows[0])==3 else [WIDTH/len(rows[0])]*len(rows[0])
    if rows[0][0]=='Time':widths=[1.02,3.5,2.42]
    if rows[0][0]=='Artifact family':widths=[1.25,2.75,2.94]
    for c,w in zip(table.columns,widths):c.width=Inches(w)
    pr=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
        borders.append(e)
    pr.append(borders)
    for ri,items in enumerate(rows):
        row=table.rows[0] if ri==0 else table.add_row()
        rp=row._tr.get_or_add_trPr()
        rp.append(OxmlElement('w:cantSplit'))
        if ri==0:rp.append(OxmlElement('w:tblHeader'))
        for ci,text in enumerate(items):
            cell=row.cells[ci];cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side,val in [('top','90'),('bottom','90'),('left','100'),('right','100')]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');margins.append(e)
            cp.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'363B40' if ri==0 else 'F4F5F6' if ri%2==0 else 'FFFFFF');cp.append(shade)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.08
            inline(p,text,table_head=ri==0)
            for r in p.runs:r.font.size=Pt(9.2)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def numbered_paragraph(text, number):
    # Real numbering with a separate concrete definition per reading route.
    if number==1:
        root=doc.part.numbering_part.element
        aid=max([int(e.get(qn('w:abstractNumId'))) for e in root.findall(qn('w:abstractNum'))] or [0])+1
        nid=max([int(e.get(qn('w:numId'))) for e in root.findall(qn('w:num'))] or [0])+1
        abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),str(aid))
        lvl=OxmlElement('w:lvl');lvl.set(qn('w:ilvl'),'0')
        for tag,val in [('start','1'),('numFmt','decimal'),('lvlText','%1.')]:
            item=OxmlElement('w:'+tag);item.set(qn('w:val'),val);lvl.append(item)
        pp=OxmlElement('w:pPr');ind=OxmlElement('w:ind');ind.set(qn('w:left'),'300');ind.set(qn('w:hanging'),'240');pp.append(ind);lvl.append(pp)
        abstract.append(lvl);root.append(abstract)
        num=OxmlElement('w:num');num.set(qn('w:numId'),str(nid));ref=OxmlElement('w:abstractNumId');ref.set(qn('w:val'),str(aid));num.append(ref);root.append(num)
        numbered_paragraph.num_id=nid
    p=doc.add_paragraph()
    pp=p._p.get_or_add_pPr();n=OxmlElement('w:numPr')
    level=OxmlElement('w:ilvl');level.set(qn('w:val'),'0');n.append(level)
    val=OxmlElement('w:numId');val.set(qn('w:val'),str(numbered_paragraph.num_id));n.append(val);pp.append(n)
    inline(p,text)

lines=SOURCE.read_text().splitlines()
headings=[line[3:] for line in lines if line.startswith('## ')]
appendix=False
appendix_group=0
frontmatter=True
i=0
while i<len(lines):
    line=lines[i]
    if not line.strip():i+=1;continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:],'Title');i+=1;continue
    if line.startswith('## '):
        text=line[3:]
        if frontmatter:
            doc.add_paragraph('Contents','Heading 1')
            for no,heading in enumerate(headings,1):
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.line_spacing=1.0
                hyperlink(p,heading,f'section_{no}',True)
            doc.add_page_break();frontmatter=False
        p=doc.add_paragraph(text,'Heading 1')
        no=int(text.split()[0]);bookmark(p,f'section_{no}')
        if no==21:p.paragraph_format.page_break_before=True
        appendix=no==21
        i+=1;continue
    if line.startswith('### '):
        p=doc.add_paragraph(line[4:],'Heading 2')
        if appendix:
            if appendix_group:p.paragraph_format.page_break_before=True
            appendix_group+=1
        i+=1;continue
    if line.startswith('!['):
        match=re.match(r'!\[([^\]]*)\]\((.*)\)',line)
        caption,path=match.groups();path=(HERE/path).resolve()
        with Image.open(path) as im:w,h=im.size
        picwidth=min(WIDTH, 2.78*w/h)
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_after=Pt(4)
        r=p.add_run();r.add_picture(str(path),width=Inches(picwidth))
        cap=doc.add_paragraph(style='Caption');inline(cap,caption)
        cap.paragraph_format.keep_with_next=True
        i+=1;continue
    if line.startswith('|'):
        batch=[]
        while i<len(lines) and lines[i].startswith('|'):batch.append(lines[i]);i+=1
        add_table(batch);continue
    if line.startswith('```'):
        i+=1
        while i<len(lines) and not lines[i].startswith('```'):
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2)
            p.paragraph_format.line_spacing=1.0
            r=p.add_run(lines[i]);r.font.name='Courier New';r.font.size=Pt(8.5)
            i+=1
        i+=1;continue
    match=re.match(r'^(\d+)\. (.*)$',line)
    if match:
        numbered_paragraph(match.group(2),int(match.group(1)));i+=1;continue
    if line.startswith('- '):
        p=doc.add_paragraph(style='List Bullet');inline(p,line[2:]);i+=1;continue
    chunk=[line];i+=1
    while i<len(lines) and lines[i].strip() and not re.match(r'^(#|!\[|\||```|\d+\. |\- )',lines[i]):chunk.append(lines[i]);i+=1
    p=doc.add_paragraph();inline(p,' '.join(chunk))
    if frontmatter:
        for r in p.runs:r.font.size=Pt(9.5)
        p.paragraph_format.space_after=Pt(7)

# Navigation in Word and accessibility metadata for the original visual plates.
for drawing,image_line in zip(doc.element.findall('.//'+qn('wp:docPr')), [x for x in lines if x.startswith('![')]):
    drawing.set('descr',re.match(r'!\[([^\]]*)\]',image_line).group(1))
settings=doc.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
doc.save(OUTPUT)
print(OUTPUT)
print(f'{len(doc.paragraphs)} paragraphs; {len(doc.tables)} tables; {len(doc.inline_shapes)} images')
